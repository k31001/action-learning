"""시각 우선 덱 도구 (samsung-memory-ppt-design-skill v2, 11절).

20 x 11.25in 캔버스, python-pptx. 레퍼런스 구현: 다운턴에서 배운 SSD 생존 전략 3장 덱
(k31001/action-learning `outputs/presentation/scripts/generate_ssd_survival_pptx.py`).

사용 예:
    from deck_helpers import Deck, BLUE, GRAY, INK, WHITE
    d = Deck("덱 이름", total=3, logos_dir="assets/logos", photos_dir="assets/photos", renders_dir="assets/renders")
    s = d.slide(1, "배경", "앞 절의 결론,\\n뒤 절의 결론")
    d.panel_head(s, 0.79, 5.6, 1, "교훈: 한 줄 결론")
    d.fit(s, d.part("ssd"), 6.8, 3.0, 2.0, 1.0)
    d.band(s, 9.60, 0.80, "명제", "한 줄 명제")
    d.footer(s, "출처: ...")
    d.save("out.pptx")
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---- 토큰 (2.A)
BLUE = RGBColor(0x14, 0x28, 0xA0)
BLUE_T1 = RGBColor(0x3C, 0x5A, 0xC8)
BLUE_T2 = RGBColor(0xAA, 0xB8, 0xE8)
TINT = RGBColor(0xEE, 0xF1, 0xFA)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
GRAY_2 = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xD9, 0xD9, 0xD9)
PALE = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ---- 캔버스 (11.B)
MX, CW = 0.79, 18.42
RIGHT = MX + CW
SIZE = {"title": 32, "panel": 24, "body": 18, "caption": 20, "big": 48, "band": 24, "source": 15, "top": 18}


class Deck:
    def __init__(self, name, total, grade="[문서등급 표기]", font=None, font_ea=None,
                 logos_dir="assets/logos", photos_dir="assets/photos", renders_dir="assets/renders"):
        self.name, self.total, self.grade = name, total, grade
        self.font = font or os.environ.get("FONT_LATIN", "Arial")
        self.font_ea = font_ea or os.environ.get("FONT_EA", self.font)
        self.logos_dir, self.photos_dir, self.renders_dir = logos_dir, photos_dir, renders_dir
        self.prs = Presentation()
        self.prs.slide_width = Emu(18288000)   # 20.00 in
        self.prs.slide_height = Emu(10287000)  # 11.25 in
        self.no = 0

    # ------------------------------------------------ 기본 도형
    def _font(self, run, size, bold, color):
        f = run.font
        f.name, f.size, f.bold = self.font, Pt(size), bold
        f.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = rPr.makeelement(qn("a:ea"), {})
            rPr.append(ea)
        ea.set("typeface", self.font_ea)

    def tb(self, s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
        """paras: [(text, size, bold, color)] 또는 [[run, run], ...]. 항목당 문단 1개."""
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            runs = para if isinstance(para[0], (list, tuple)) else [para]
            for (t, size, bold, color) in runs:
                r = p.add_run()
                r.text = t
                self._font(r, size, bold, color)
        return box

    def rect(self, s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, dash=False):
        sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.shadow.inherit = False
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line
            sp.line.width = Pt(lw)
            if dash:
                sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, 0.08 / max(0.01, min(w, h)))
        return sp

    def label_box(self, s, x, y, w, h, paras, fill=None, line=None, lw=1.0, align=PP_ALIGN.CENTER, dash=False):
        self.rect(s, x, y, w, h, fill=fill, line=line, lw=lw, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=dash)
        pad = 0.14 if align != PP_ALIGN.CENTER else 0.06
        self.tb(s, x + pad, y, w - 2 * pad, h, paras, align=align, anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------ 이미지·로고 (11.F, 11.G)
    def img(self, s, path, x, y, w=None, h=None):
        with Image.open(path) as im:
            a = im.size[0] / im.size[1]
        w = h * a if w is None else w
        h = w / a if h is None else h
        s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        return w, h

    def fit(self, s, path, x, y, bw, bh, align="center"):
        """박스 안에 비율 유지로 맞춘다. (x, y, w, h) 반환."""
        with Image.open(path) as im:
            a = im.size[0] / im.size[1]
        w, h = (bw, bw / a) if bw / a <= bh else (bh * a, bh)
        ox = x + (bw - w) / 2 if align == "center" else x
        oy = y + (bh - h) / 2
        s.shapes.add_picture(path, Inches(ox), Inches(oy), Inches(w), Inches(h))
        return ox, oy, w, h

    def part(self, name):
        """부품 이미지: 사진 폴더 우선, 없으면 3D 렌더(render_<name>.png)."""
        for ext in (".png", ".jpg", ".jpeg"):
            p = os.path.join(self.photos_dir, name + ext)
            if os.path.exists(p):
                return p
        return os.path.join(self.renders_dir, f"render_{name}.png")

    def logo(self, name):
        return os.path.join(self.logos_dir, f"{name}.png")

    def chip(self, s, text, x, y, h, size=18):
        """로고를 못 구한 기업의 워드마크 칩. 폭 반환."""
        w = 0.40 + 0.16 * size / 18 * len(text)
        self.rect(s, x, y, w, h, fill=WHITE, line=GRAY_2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        self.tb(s, x, y, w, h, [(text, size, True, INK)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return w

    # ------------------------------------------------ 골격 (11.C)
    def slide(self, no, section, title):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.no = no
        self.tb(s, MX, 0.40, 12.0, 0.36, [[(self.name, 18, True, BLUE), (f"   {no} / {self.total}  {section}", 18, False, GRAY)]])
        self.tb(s, RIGHT - 6.0, 0.40, 6.0, 0.36, [(self.grade, 18, False, GRAY)], align=PP_ALIGN.RIGHT)
        self.tb(s, MX, 0.86, CW, 1.12, [(ln, SIZE["title"], True, INK) for ln in title.split("\n")],
                anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
        self.rect(s, MX, 2.12, CW, 0.02, fill=LINE)
        return s

    def footer(self, s, text):
        self.tb(s, MX, 10.50, 16.9, 0.62, [(text, SIZE["source"], False, GRAY)], spacing=1.05)
        self.tb(s, RIGHT - 1.2, 10.50, 1.2, 0.34, [(f"{self.no} / {self.total}", 16, False, GRAY)], align=PP_ALIGN.RIGHT)

    def band(self, s, y, h, tag, text, size=24, pills=None):
        """Blue 결론 밴드. pills=['기술','인재'] 이면 오른쪽에 다음 장 알약."""
        self.rect(s, MX, y, CW, h, fill=BLUE)
        self.tb(s, MX + 0.40, y, 1.35, h, [(tag, 20, False, BLUE_T2)], anchor=MSO_ANCHOR.MIDDLE)
        pw = sum(0.40 + 0.28 * len(p) for p in pills) + 0.18 * (len(pills) - 1) + 0.35 if pills else 0.0
        self.tb(s, MX + 1.75, y, CW - 1.75 - pw - 0.2, h, [(text, size, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        cx = RIGHT - pw + 0.05
        for p in pills or []:
            w = 0.40 + 0.28 * len(p)
            self.rect(s, cx, y + (h - 0.52) / 2, w, 0.52, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            self.tb(s, cx, y + (h - 0.52) / 2, w, 0.52, [(p, 20, True, BLUE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cx += w + 0.18

    def panel_head(self, s, x, w, no, text, y=2.34):
        d = 0.48
        self.rect(s, x, y, d, d, fill=BLUE, shape=MSO_SHAPE.OVAL)
        self.tb(s, x, y, d, d, [(str(no), 20, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.tb(s, x + d + 0.16, y - 0.04, w - d - 0.16, d + 0.08, [(text, SIZE["panel"], True, INK)], anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------ 개념 도형 어휘 (11.E)
    def chevron(self, s, x, y, w=0.18, h=0.52, fill=BLUE_T2):
        self.rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.CHEVRON)

    def arrow_r(self, s, x, y, w, h, fill=GRAY_2):
        return self.rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW)

    def down(self, s, cx, y, fill=BLUE_T2):
        """지금 → 앞으로 사이 아래 삼각형."""
        self.rect(s, cx - 0.22, y, 0.44, 0.26, fill=fill, shape=MSO_SHAPE.ISOSCELES_TRIANGLE).rotation = 180

    def bracket(self, s, x, y, w, color, label, size=18):
        """위쪽 범위 괄호 + 라벨."""
        self.rect(s, x, y + 0.36, w, 0.035, fill=color)
        self.rect(s, x, y + 0.36, 0.035, 0.16, fill=color)
        self.rect(s, x + w - 0.035, y + 0.36, 0.035, 0.16, fill=color)
        self.tb(s, x, y - 0.02, w, 0.36, [(label, size, True, color)], align=PP_ALIGN.CENTER)

    def person(self, s, x, y, h, color=BLUE):
        """사람 아이콘(머리 원 + 몸통). 폭(0.62h) 반환."""
        w, d = 0.62 * h, 0.36 * h
        self.rect(s, x + (w - d) / 2, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
        b = self.rect(s, x, y + 0.42 * h, w, 0.58 * h, fill=color, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        b.adjustments[0] = 0.45
        return w

    def chain(self, s, x, y, w, h, texts, gap=0.40, hot_last=True):
        """원인 → 결과 체인. 마지막 상자만 Blue."""
        for i, t in enumerate(texts):
            bx = x + i * (w + gap)
            last = hot_last and i == len(texts) - 1
            self.label_box(s, bx, y, w, h, [(ln, 18, True, WHITE if last else INK) for ln in t.split("\n")],
                           fill=BLUE if last else PALE)
            if i < len(texts) - 1:
                self.arrow_r(s, bx + w + 0.07, y + h / 2 - 0.13, gap - 0.14, 0.26)

    def big_number(self, s, x, y, w, pre, num, post, sub=None):
        self.tb(s, x, y, w, 0.80, [[(pre, 26, True, BLUE), (num, 48, True, BLUE), (post, 26, True, BLUE)]], anchor=MSO_ANCHOR.MIDDLE)
        if sub:
            self.tb(s, x, y + 0.78, w, 0.36, [(sub, 18, False, GRAY)])

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.prs.save(path)
        return path
