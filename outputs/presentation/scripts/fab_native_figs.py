"""시안 1~6 기둥 그림 9종 x 색 테마 2종을 PowerPoint 기본 도형으로 그린다 (v2, 2026-10-15).

PNG(generate_fab_common_platform_figs.py)와 같은 그림을 사각형 · 원 · 선 · 화살표 · 자유형 다각형 · 텍스트 상자로 다시 그려,
슬라이드에서 바로 고칠 수 있게 한다. 그림마다 그룹 1개(그룹 해제하면 도형 하나씩 편집).
  - 아이소메트릭 블록: 면 3개(왼쪽 · 오른쪽 · 윗면)를 자유형 다각형으로, 블록 하나가 하위 그룹 1개
  - 수요 곡선 · 누적 면적: 자유형(꼭짓점 편집 가능)
좌표는 PNG와 같은 그림 좌표(폭 5.30 x 높이 1.80in, 왼쪽 아래 원점, y 위로)로 쓰고 슬라이드 좌표로 바꾼다.
색 테마는 PNG 생성기의 MONO / COLOR를 그대로 쓴다(단일 소스).
"""
import math

import numpy as np
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from generate_fab_common_platform_figs import COLOR, HUES, MONO, INK, G, G2, MID, LINE, WHITE, H, mix

C30 = math.cos(math.radians(30))


def rgb(hexc):
    return RGBColor.from_string(hexc.lstrip("#"))


def theme(style, effect):
    """style 'v*' = 단일 Blue, 'c*' = 다색(효과별 색)."""
    if style.startswith("v"):
        return dict(MONO)
    K = dict(COLOR)
    h, lt, bg = HUES[effect]
    K.update(hue=h, hue_lt=lt, hue_bg=bg, arrow=h, arrow_lt=lt)
    return K


def text_w(s, size):
    k = size / 16
    return sum(0.235 * k if ord(ch) > 0x2E80 else (0.07 * k if ch == " " else 0.125 * k) for ch in s) + 0.08


class Canvas:
    """그림 좌표(인치, y 위로) → 슬라이드 좌표. 도형은 그룹 하나에 담는다."""

    def __init__(self, d, parent_shapes, ox, oy, name):
        self.d, self.ox, self.oy = d, ox, oy
        self.g = parent_shapes.add_group_shape()
        self.g.name = name
        self.sh = self.g.shapes

    def sub(self, name):
        """하위 그룹(블록 하나 등)."""
        c = Canvas.__new__(Canvas)
        c.d, c.ox, c.oy = self.d, self.ox, self.oy
        c.g = self.sh.add_group_shape()
        c.g.name = name
        c.sh = c.g.shapes
        return c

    def X(self, x):
        return Inches(self.ox + x)

    def Y(self, y):
        return Inches(self.oy + H - y)

    # ---------------------------------------------------------- 스타일
    @staticmethod
    def _style(sp, fc, ec, lw, dash, hatch=False):
        sp.shadow.inherit = False
        if hatch:
            sp.fill.patterned()
            sp.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
            sp.fill.fore_color.rgb = rgb(MID)
            sp.fill.back_color.rgb = rgb(fc or WHITE)
        elif fc is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fc)
        if ec is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = rgb(ec)
            sp.line.width = Pt(lw)
            if dash:
                sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return sp

    # ---------------------------------------------------------- 도형
    def shape(self, kind, x, y, w, h, fc=None, ec=None, lw=0.8, dash=False, r=None, hatch=False):
        """(x, y) = 왼쪽 아래."""
        sp = self.sh.add_shape(kind, self.X(x), self.Y(y + h), Inches(w), Inches(h))
        if r is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r / max(0.01, min(w, h)))
        return self._style(sp, fc, ec, lw, dash, hatch)

    def rect(self, x, y, w, h, fc=None, ec=None, lw=0.8, dash=False, hatch=False):
        return self.shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fc, ec, lw, dash, hatch=hatch)

    def rrect(self, x, y, w, h, r, fc=None, ec=None, lw=0.8, dash=False):
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fc, ec, lw, dash, r=r)

    def hatch(self, x, y, w, h):
        return self.rect(x, y, w, h, fc=WHITE, ec=MID, lw=0.8, dash=True, hatch=True)

    def circle(self, cx, cy, r, fc, ec=WHITE, lw=1.0):
        return self.shape(MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fc, ec, lw)

    def ellipse(self, cx, cy, w, h, fc, ec=WHITE, lw=1.0):
        return self.shape(MSO_SHAPE.OVAL, cx - w / 2, cy - h / 2, w, h, fc, ec, lw)

    def poly(self, pts, fc=None, ec=None, lw=0.8, close=True, dash=False, hatch=False):
        """자유형 다각형(close=False 이면 열린 선)."""
        e = [(int(self.X(x)), int(self.Y(y))) for (x, y) in pts]
        b = self.sh.build_freeform(e[0][0], e[0][1], scale=1.0)
        b.add_line_segments(e[1:], close=close)
        sp = b.convert_to_shape()
        return self._style(sp, fc if close else None, ec, lw, dash, hatch)

    def line(self, x1, y1, x2, y2, color, lw=1.0, head=False, tail=False, dash=False, dot=False):
        c = self.sh.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(x1), self.Y(y1), self.X(x2), self.Y(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(lw)
        if dash:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = c.line._get_or_add_ln()
        if tail:
            ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        if head:
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        return c

    def block_arrow(self, x, y, w, h, color):
        """PNG와 같은 굵은 화살표. (x, y) = 왼쪽 끝 · 세로 가운데."""
        return self.poly([(x, y - h * 0.22), (x + w * 0.58, y - h * 0.22), (x + w * 0.58, y - h / 2), (x + w, y),
                          (x + w * 0.58, y + h / 2), (x + w * 0.58, y + h * 0.22), (x, y + h * 0.22)], fc=color)

    def swap(self, cx, y, w=0.26, h=0.12):
        """설비 전환 표시: 위로 휜 화살표. (cx, y) = 가운데 · 아래."""
        return self.shape(MSO_SHAPE.CURVED_DOWN_ARROW, cx - w / 2, y, w, h, fc=INK)

    def text(self, x, y, s, size=16, color=G, bold=False, ha="left", w=None, rot=0):
        """(x, y) = 기준점(세로 가운데). ha: left · center · right."""
        w = w or text_w(s, size)
        h = 0.30 * size / 16
        bx = x if ha == "left" else (x - w / 2 if ha == "center" else x - w)
        tb = self.sh.add_textbox(self.X(bx), self.Y(y + h / 2), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = s.split("\n")
        for i, ln_ in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[ha]
            r = p.add_run()
            r.text = ln_
            self.d._font(r, size, bold, rgb(color))
        if rot:
            tb.rotation = rot
        return tb


# ------------------------------------------------------------------ 아이소메트릭
class Iso:
    def __init__(self, c, ox, oy, s):
        self.c, self.ox, self.oy, self.s = c, ox, oy, s

    def p(self, x, y, z=0.0):
        return (self.ox + (x - y) * C30 * self.s, self.oy + (x + y) * 0.5 * self.s + z * self.s)

    def plate(self, n, m, fc=WHITE, ec=MID, lw=0.8, grid=True):
        c = self.c
        c.poly([self.p(0, 0), self.p(n, 0), self.p(n, m), self.p(0, m)], fc=fc, ec=ec, lw=lw)
        if grid:
            for i in range(1, n):
                c.line(*self.p(i, 0), *self.p(i, m), LINE, 0.6)
            for j in range(1, m):
                c.line(*self.p(0, j), *self.p(n, j), LINE, 0.6)

    def tile(self, x, y, w=1, d=1, fc=None, ec=MID, hatch=False, dash=False, lw=0.9):
        p = self.p
        self.c.poly([p(x, y), p(x + w, y), p(x + w, y + d), p(x, y + d)], fc=fc, ec=ec, lw=lw, dash=dash, hatch=hatch)

    def box(self, x, y, w, d, h, color, z0=0.0, top=None, side=None, ec=WHITE, lw=0.6, name="설비"):
        p = self.p
        topc = top or mix(color, 0.35)
        lc = side or color
        rc = mix(side, -0.18) if side else mix(color, -0.22)
        g = self.c.sub(name)
        g.poly([p(x, y, z0), p(x, y + d, z0), p(x, y + d, z0 + h), p(x, y, z0 + h)], fc=lc, ec=ec, lw=lw)
        g.poly([p(x, y, z0), p(x + w, y, z0), p(x + w, y, z0 + h), p(x, y, z0 + h)], fc=rc, ec=ec, lw=lw)
        g.poly([p(x, y, z0 + h), p(x + w, y, z0 + h), p(x + w, y + d, z0 + h), p(x, y + d, z0 + h)], fc=topc, ec=ec, lw=lw)


def far_first(items):
    return sorted(items, key=lambda t: -(t[0] + t[1]))


# ================================================================== 시안 1 · 4 아이소메트릭
def iso_space(c, K):
    s = 0.27
    L = Iso(c, 1.17, 0.10, s)
    L.plate(4, 4)
    for (x, y) in [(0, 2), (1, 0), (2, 3), (3, 1), (1, 3)]:
        L.tile(x + 0.06, y + 0.06, 0.88, 0.88, fc=WHITE, ec=MID, hatch=True, dash=True)
    c.line(*L.p(2, 0), *L.p(2, 4), G, 3.0)                      # DRAM | NAND 벽
    a, b, cc = K["mixA"], K["mixB"], K["mixC"]
    tools = [(0, 0, 0.8, 0.7, 1.0, a), (0, 1, 0.6, 0.8, 0.7, b), (0, 3, 0.7, 0.6, 1.25, cc), (1, 1, 0.8, 0.8, 0.9, a),
             (1, 2, 0.6, 0.6, 0.6, b), (2, 0, 0.7, 0.8, 1.1, cc), (2, 1, 0.8, 0.6, 0.75, b), (2, 2, 0.6, 0.7, 1.3, a),
             (3, 0, 0.6, 0.6, 0.65, b), (3, 2, 0.8, 0.8, 0.95, cc), (3, 3, 0.7, 0.7, 0.8, a)]
    for (x, y, w, d, h, col) in far_first(tools):
        L.box(x + (1 - w) / 2, y + (1 - d) / 2, w, d, h, col)
    c.text(0.02, 1.66, "Fab 분리", 16, G, bold=True)
    c.text(0.02, 1.40, "DRAM | NAND", 16, G2)
    c.block_arrow(2.30, 0.72, 0.50, 0.30, K["arrow"])
    R = Iso(c, 4.03, 0.10, s)
    R.plate(4, 4, ec=K["hue_lt"])
    for yy in range(4):
        R.tile(3.06, yy + 0.06, 0.88, 0.88, fc=K["gain_bg"], ec=K["gain"], dash=True, lw=1.1)
    rows = {0: (0.80, K["modA"]), 1: (0.80, K["modA"]), 2: (1.10, K["modB"]), 3: (0.95, K["modC"])}
    for (x, y) in far_first([(x, y) for x in range(3) for y in range(4)]):
        h, col = rows[y]
        R.box(x + 0.08, y + 0.08, 0.84, 0.84, h, col)
    c.text(3.02, 1.66, "Common Platform", 16, K["hue"], bold=True)
    px, py = R.p(4.0, 2.0)
    c.line(px + 0.02, py + 0.02, px + 0.20, py + 0.20, K["gain"], 1.0)
    c.text(px + 0.22, py + 0.28, "확보", 16, K["gain_txt"], bold=True)


def iso_agility(c, K):
    s = 0.31
    F = Iso(c, 0.66, 0.12, s)
    F.plate(6, 2, ec=K["hue_lt"])
    plan = {(0, 0): "D", (1, 0): "D", (2, 0): "S", (3, 0): "N", (4, 0): "N", (5, 0): "N",
            (0, 1): "D", (1, 1): "D", (2, 1): "D", (3, 1): "S", (4, 1): "N", (5, 1): "N"}
    for (x, y) in far_first(list(plan)):
        k = plan[(x, y)]
        if k == "D":
            F.box(x + 0.1, y + 0.1, 0.8, 0.8, 1.0, K["dram"], name="DRAM 설비")
        elif k == "N":
            F.box(x + 0.1, y + 0.1, 0.8, 0.8, 1.0, K["nand"], name="NAND 설비")
        else:
            F.box(x + 0.1, y + 0.1, 0.8, 0.8, 1.0, K["nand"], top=K["dram_lt"], side=K["nand"], name="전환 중 설비")
    for (x, y) in ((2, 0), (3, 1)):
        cx, cy = F.p(x + 0.5, y + 0.5, 1.0)
        c.swap(cx, cy + 0.06, 0.32, 0.14)
    c.text(0.02, 1.66, "공용 설비 전환", 16, INK, bold=True)
    x0, x1 = 2.78, 4.52
    t = np.linspace(0, 1, 49)
    dr = 1.30 + 0.24 * np.sin(2 * np.pi * (t * 1.25 - 0.10))
    na = 1.30 - 0.24 * np.sin(2 * np.pi * (t * 1.25 - 0.10))
    X = x0 + t * (x1 - x0)
    c.line(x0, 0.90, x1, 0.90, LINE, 0.8)
    nseg = 6
    sw = (x1 - x0) / nseg
    for i in range(nseg):
        tm = (i + 0.5) / nseg
        share = 0.5 + 0.38 * math.sin(2 * np.pi * (tm * 1.25 - 0.10))
        hh = 0.50
        c.rect(x0 + i * sw + 0.02, 0.18, sw - 0.04, hh * share, fc=K["dram"])
        c.rect(x0 + i * sw + 0.02, 0.18 + hh * share, sw - 0.04, hh * (1 - share), fc=K["nand_lt"])
    c.poly(list(zip(X, dr)), ec=K["dram"], lw=2.6, close=False)
    c.poly(list(zip(X, na)), ec=K["nand"], lw=2.2, close=False)
    c.text(x1 + 0.07, dr[-1], "DRAM", 16, K["dram"], bold=True)
    c.text(x1 + 0.07, na[-1], "NAND", 16, K["nand_txt"])
    c.text(x0, 1.70, "수요", 16, G)
    c.text(x1 + 0.07, 0.43, "배분", 16, G)


def iso_unif(c, K):
    s = 0.30
    L = Iso(c, 0.82, 0.14, s)
    L.plate(3, 3)
    a, b, cc = K["mixA"], K["mixB"], K["mixC"]
    mixed = [(0, 0, 0.8, 0.6, 1.0, a), (1, 0, 0.6, 0.8, 0.6, b), (2, 0, 0.7, 0.7, 1.3, cc),
             (0, 1, 0.6, 0.6, 1.4, cc), (1, 1, 0.8, 0.8, 0.8, a), (2, 1, 0.6, 0.6, 0.7, b),
             (0, 2, 0.7, 0.8, 0.75, b), (1, 2, 0.6, 0.6, 1.15, a), (2, 2, 0.8, 0.7, 0.95, cc)]
    for (x, y, w, d, h, col) in far_first(mixed):
        L.box(x + (1 - w) / 2, y + (1 - d) / 2, w, d, h, col)
    c.text(0.02, 1.66, "라인별 모델 혼재", 16, G, bold=True)
    Iso(c, 2.05, 0.16, 0.24).box(0, 0, 1, 1, 0.9, MID, name="발주 1단")
    c.text(2.05, 0.80, "발주", 16, G, ha="center")
    c.block_arrow(2.42, 0.70, 0.46, 0.28, K["arrow"])
    R = Iso(c, 3.48, 0.14, s)
    R.plate(3, 2, ec=K["hue_lt"])
    for (x, y) in far_first([(x, y) for x in range(3) for y in range(2)]):
        R.box(x + 0.1, y + 0.1, 0.8, 0.8, 1.05, K["com"], name="공용 모델")
    c.text(2.98, 1.66, "공용 모델", 16, K["com"], bold=True)
    T2 = Iso(c, 4.86, 0.16, 0.24)
    for k in range(4):
        T2.box(0, 0, 1, 1, 0.9, K["gain"], z0=k * 0.95, name=f"발주 {k + 1}단")
    c.text(4.86, 1.52, "발주 ↑", 16, K["gain_txt"], bold=True, ha="center")


# ================================================================== 시안 2 · 5 메커니즘
def mech_space(c, K):
    pal = {"t": dict(fc=K["mixA"]), "T": dict(fc=K["modA"]), "a": dict(fc=K["add"]),
           "d": dict(fc=WHITE, ec=MID, dash=True, hatch=True), "o": dict(fc="#E2E2E2")}
    before = ["tdttdtdt", "dttdttdd", "oooooooo", "ttdtdttd"]
    after = ["TTTTTTaa", "TTTTTTad", "oooooooo", "TTTTTTaa"]
    cell, gap = 0.22, 0.04

    def waffle(x0, y0, plan, name):
        g = c.sub(name)
        for r, row in enumerate(plan):
            for k, ch in enumerate(row):
                g.rect(x0 + k * (cell + gap), y0 + (len(plan) - 1 - r) * (cell + gap), cell, cell, **pal[ch])

    c.text(0.06, 1.66, "Fab 분리", 16, G, bold=True)
    c.text(3.12, 1.66, "Common Platform", 16, K["hue"], bold=True)
    waffle(0.06, 0.46, before, "분리 격자")
    c.block_arrow(2.24, 0.85, 0.62, 0.30, K["arrow"])
    waffle(3.12, 0.46, after, "통합 격자")
    c.rect(3.09, 0.43, 8 * 0.26 + 0.02, 4 * 0.26 + 0.02, fc=None, ec=K["hue"], lw=1.2)
    lx = 0.06
    for lab, ch in (("설비", "T"), ("추가 설비", "a"), ("Dead Space", "d"), ("통로", "o")):
        c.rect(lx, 0.10, 0.18, 0.18, **pal[ch])
        c.text(lx + 0.25, 0.19, lab, 16, G)
        lx += 0.25 + {"설비": 0.50, "추가 설비": 0.98, "Dead Space": 1.32, "통로": 0.5}[lab]


def mech_agility(c, K):
    x0, x1, yb = 0.05, 3.05, 0.12
    t = np.linspace(0, 1, 91)
    dram = 0.52 + 0.20 * np.sin(2 * np.pi * t * 1.5)
    nand = 0.48 + 0.17 * np.sin(2 * np.pi * t * 1.5 + 0.9 * np.pi)
    tt = np.linspace(0, 1, 601)                                   # 피크는 촘촘한 표본으로
    d_ = 0.52 + 0.20 * np.sin(2 * np.pi * tt * 1.5)
    n_ = 0.48 + 0.17 * np.sin(2 * np.pi * tt * 1.5 + 0.9 * np.pi)
    X = x0 + t * (x1 - x0)
    top_d = yb + dram
    top_n = top_d + nand
    c.poly([(x0, yb)] + list(zip(X, top_d)) + [(x1, yb)], fc=K["dram"])
    c.poly(list(zip(X, top_d)) + list(zip(X[::-1], top_n[::-1])), fc=K["nand_lt"])
    sep = yb + d_.max() + n_.max()
    com = yb + (d_ + n_).max()
    c.line(x0, sep, x1, sep, G2, 2.0, dash=True)
    c.line(x0, com, x1, com, K["line_ok"], 2.4)
    c.line(x0, yb, x1, yb, LINE, 0.8)
    c.text(x0 + 0.02, sep + 0.13, "분리 시 필요 설비", 16, G)
    c.text(x0 + 0.02, com + 0.12, "통합 시 필요 설비", 16, K["line_ok"], bold=True)
    c.line(x1 - 0.12, sep - 0.01, x1 - 0.12, com + 0.02, K["idle_txt"], 1.6, head=True, tail=True)
    c.text(x1 - 0.22, (sep + com) / 2 + 0.01, "유휴 ↓", 16, K["idle_txt"], bold=True, ha="right")
    c.text(x0 + 0.40, yb + 0.22, "DRAM", 16, WHITE, bold=True)
    c.text(x0 + 1.55, yb + 0.88, "NAND", 16, K["nand_lbl"], bold=True)
    rx = 3.30
    c.line(rx - 0.12, 0.08, rx - 0.12, 1.74, LINE, 0.8)
    for g_, gy in (("설비 전환 시간", 1.60), ("Bottleneck 대기", 0.84)):
        c.text(rx, gy, g_, 16, INK, bold=True)
        for k, (lab, ln, col) in enumerate((("분리", 1.30, G2), ("통합", 0.45, K["tobe"]))):
            yy = gy - 0.30 - k * 0.24
            c.text(rx, yy, lab, 16, G if k == 0 else K["tobe_txt"], bold=k == 1)
            c.rect(rx + 0.52, yy - 0.08, ln, 0.16, fc=col)


def mech_unif(c, K):
    base = 0.30
    c.text(0.05, 1.66, "설비 모델 수", 16, INK, bold=True)
    bw, u = 0.52, 0.13
    xa, xb = 0.92, 1.66
    for k in range(8):
        c.rect(xa, base + k * u, bw, u - 0.02, fc=K["dram_lt"] if k < 4 else K["nand_lt"])
    for k in range(3):
        c.rect(xb, base + k * u, bw, u - 0.02, fc=K["com"])
    c.text(xa - 0.08, base + 2 * u, "DRAM", 16, K["dram_lt"], bold=True, ha="right")
    c.text(xa - 0.08, base + 6 * u, "NAND", 16, K["nand_txt"], ha="right")
    c.text(xb + bw / 2, base + 3 * u + 0.16, "↓", 18, K["com"], bold=True, ha="center")
    c.line(0.15, base, 2.15, base, LINE, 0.8)
    c.text(xa + bw / 2, base - 0.15, "분리", 16, G, ha="center")
    c.text(xb + bw / 2, base - 0.15, "통합", 16, K["com"], bold=True, ha="center")
    c.block_arrow(2.32, 0.80, 0.42, 0.26, K["arrow_lt"])
    ox = 2.95
    c.text(ox, 1.66, "모델당 발주 물량", 16, INK, bold=True)
    xa, xb = ox + 0.15, ox + 1.05
    c.rect(xa, base, bw, 0.30, fc=G2)
    c.rect(xb, base, bw, 1.00, fc=K["com"])
    for k in range(1, 4):
        c.line(xb, base + k * 0.25, xb + bw, base + k * 0.25, WHITE, 0.9)
    c.line(ox, base, ox + 2.25, base, LINE, 0.8)
    c.text(xa + bw / 2, base - 0.15, "분리", 16, G, ha="center")
    c.text(xb + bw / 2, base - 0.15, "통합", 16, K["com"], bold=True, ha="center")
    c.text(xb + bw + 0.08, base + 0.82, "구매", 16, K["gain_txt"], bold=True)
    c.text(xb + bw + 0.08, base + 0.56, "협상력 ↑", 16, K["gain_txt"], bold=True)


# ================================================================== 시안 3 · 6 Fab 도면
def plan_space(c, K):
    fy, fh, fw = 0.06, 1.42, 2.46
    lx = 0.04
    c.rect(lx, fy, fw, fh, fc=WHITE, ec=G2, lw=1.0)
    c.line(lx + fw / 2, fy, lx + fw / 2, fy + fh, G, 2.4)
    specs = [
        (0.10, 0.12, 0.30, 0.38, "t"), (0.46, 0.12, 0.22, 0.38, "d"), (0.74, 0.12, 0.40, 0.30, "t"),
        (0.10, 0.92, 0.22, 0.44, "t"), (0.38, 0.92, 0.36, 0.44, "d"), (0.80, 0.92, 0.32, 0.36, "t"),
        (1.33, 0.12, 0.36, 0.34, "t"), (1.75, 0.12, 0.30, 0.38, "d"), (2.10, 0.12, 0.32, 0.38, "t"),
        (1.33, 0.92, 0.26, 0.40, "d"), (1.65, 0.92, 0.38, 0.44, "t"), (2.10, 0.92, 0.32, 0.30, "t")]
    for (x, y, w, h, k) in specs:
        if k == "t":
            c.rect(x, y, w, h, fc=K["zoneD"] if x < lx + fw / 2 else K["zoneN"], ec=K["zone_ec"], lw=0.8)
        else:
            c.hatch(x, y, w, h)
    for k in range(2):
        zx = lx + k * fw / 2
        c.rrect(zx + 0.10, fy + 0.60, fw / 2 - 0.20, 0.22, 0.10, fc=None, ec=G2, lw=1.6)
        for (fx, fy_) in ((zx + 0.35, fy + 0.60), (zx + 0.85, fy + 0.82)):
            c.rect(fx - 0.045, fy_ - 0.045, 0.09, 0.09, fc=G2, ec=WHITE, lw=0.6)
    c.text(lx, 1.66, "Fab 분리: OHT 2개", 16, G, bold=True)

    rx = 2.80
    c.rect(rx, fy, fw, fh, fc=WHITE, ec=K["hue"], lw=1.4)
    for (yy, col) in ((fy + 0.08, K["modA"]), (fy + 0.90, K["modB"])):
        for i in range(5):
            c.rect(rx + 0.10 + i * 0.38, yy, 0.32, 0.40, fc=col, ec=WHITE, lw=0.8)
    c.rect(rx + 2.04, fy + 0.08, 0.36, fh - 0.16, fc=K["gain_bg"], ec=K["gain"], lw=1.0, dash=True)
    c.rrect(rx + 0.10, fy + 0.60, 1.86, 0.22, 0.10, fc=None, ec=K["line_ok"], lw=2.2)
    for (fx, fy_) in ((rx + 0.40, fy + 0.60), (rx + 1.00, fy + 0.82), (rx + 1.55, fy + 0.60)):
        c.rect(fx - 0.045, fy_ - 0.045, 0.09, 0.09, fc=K["line_ok"], ec=WHITE, lw=0.6)
    c.text(rx + 2.22, fy + fh / 2, "확보", 16, K["gain_txt"], bold=True, ha="center", w=0.70, rot=270)
    c.text(rx, 1.66, "Common: OHT 1개", 16, K["hue"], bold=True)


def wafer(c, cx, cy, r, color):
    g = c.sub("웨이퍼")
    g.circle(cx, cy, r, color)
    lc = mix(color, 0.35)
    for k in (-1, 0, 1):
        g.line(cx - r * 0.75, cy + k * r * 0.38, cx + r * 0.75, cy + k * r * 0.38, lc, 0.7)
        g.line(cx + k * r * 0.38, cy - r * 0.75, cx + k * r * 0.38, cy + r * 0.75, lc, 0.7)
    g.rect(cx - r * 0.18, cy - r - 0.01, r * 0.36, 0.035, fc=WHITE)


def plan_agility(c, K):
    D, N = K["dram"], K["nand"]
    px, py, pw, ph = 1.38, 0.10, 2.20, 1.42
    c.rrect(px, py, pw, ph, 0.08, fc=K["hue_bg"], ec=K["hue"], lw=1.2, dash=True)
    for i, (sy, col) in enumerate(((1.22, D), (0.46, N))):
        for ty in (1.02, 0.40):
            c.line(0.52, sy, px - 0.02, ty + (0.08 if i == 0 else -0.06), col, 1.6, head=True)
    for (fx, fy_, col) in ((0.82, 1.17, D), (1.06, 1.07, D), (0.84, 0.50, N), (1.10, 0.47, N)):
        c.rect(fx - 0.05, fy_ - 0.05, 0.10, 0.10, fc=col, ec=WHITE, lw=0.6)
    wafer(c, 0.30, 1.22, 0.20, D)
    wafer(c, 0.30, 0.46, 0.20, N)
    c.text(0.58, 1.44, "DRAM", 16, D, bold=True)
    c.text(0.58, 0.24, "NAND", 16, K["nand_txt"], bold=True)
    c.text(px + pw / 2, 1.66, "공용 BEOL 설비 Pool", 16, K["hue"], bold=True, ha="center")
    for r_, row in enumerate(("DDSN", "DSNN")):
        for k, ch in enumerate(row):
            x, y = px + 0.16 + k * 0.50, py + 0.80 - r_ * 0.62
            c.rrect(x, y, 0.40, 0.44, 0.04, fc=D if ch == "D" else N, ec=WHITE, lw=0.8)
            if ch == "S":
                c.rect(x, y + 0.22, 0.40, 0.22, fc=D)
                c.swap(x + 0.20, y + 0.46)
    c.line(px + pw + 0.04, 0.82, 4.10, 0.82, K["hue"], 2.4, head=True)
    for k, col in enumerate(K["bond"]):
        c.ellipse(4.62, 0.62 + k * 0.16, 0.96, 0.30, col)
    c.text(4.62, 1.40, "Bonding", 16, K["hue"], bold=True, ha="center")


def plan_unif(c, K):
    steps = ["Depo", "Etch", "CMP", "Bond"]
    dram = [["A1", "B4"], ["C2", "D3"], ["E1", "F5"], ["G2", "H1"]]
    nand = [["K3", "A2"], ["L1", "C6"], ["M4", "E2"], ["N2", "G7"]]
    plat = ["P1", "P2", "P3", "P4"]
    top, rh = 1.42, 0.335
    cw, ch, pitch = 0.40, 0.25, 0.45
    xd, xn, xa, xp, xl = 0.66, 1.62, 2.58, 3.04, 3.78
    c.text(xd + (pitch + cw) / 2, 1.66, "DRAM", 16, K["chipD"][2], bold=True, ha="center")
    c.text(xn + (pitch + cw) / 2, 1.66, "NAND", 16, K["chipN"][2], bold=True, ha="center")
    c.text(xp + 0.26, 1.66, "공용", 16, K["com"], bold=True, ha="center")

    def chip(x, y, code, fc, ec, tc, bold=False):
        c.rrect(x, y - ch / 2, cw, ch, 0.04, fc=fc, ec=ec, lw=0.9)
        c.text(x + cw / 2, y, code, 16, tc, bold=bold, ha="center", w=cw)

    labels = [("레시피 공유", K["com"]), ("부품 공용", K["com"]), ("물량 집중", K["com"]), ("협상력 ↑", K["gain_txt"])]
    for i, st in enumerate(steps):
        y = top - i * rh - rh / 2
        if i:
            c.line(0.02, y + rh / 2, 5.28, y + rh / 2, LINE, 0.7)
        c.text(0.02, y, st, 16, INK, bold=True)
        for k, code in enumerate(dram[i]):
            chip(xd + k * pitch, y, code, *K["chipD"])
        for k, code in enumerate(nand[i]):
            chip(xn + k * pitch, y, code, *K["chipN"])
        c.block_arrow(xa, y, 0.36, 0.20, K["arrow_lt"])
        for k in (2, 1):                                      # 물량이 쌓인 느낌
            c.rrect(xp + k * 0.06, y - ch / 2 + k * 0.035, cw, ch, 0.04, fc=[K["com"], K["com_lt"], K["com_lt2"]][k], ec=WHITE, lw=0.8)
        chip(xp, y, plat[i], K["com"], WHITE, WHITE, bold=True)
        c.text(xl, y, labels[i][0], 16, labels[i][1], bold=True)


DRAW = {("1", "space"): iso_space, ("1", "agility"): iso_agility, ("1", "unif"): iso_unif,
        ("2", "space"): mech_space, ("2", "agility"): mech_agility, ("2", "unif"): mech_unif,
        ("3", "space"): plan_space, ("3", "agility"): plan_agility, ("3", "unif"): plan_unif}
TITLES = {"space": "그림 1 공간 효율성", "agility": "그림 2 시황 변화 대응력", "unif": "그림 3 설비 통일성"}


def draw(d, slide, style, effect, ox, oy):
    """style: 'v1'~'v3' (단일 Blue) · 'c1'~'c3' (다색), effect: space · agility · unif. 그룹 1개를 돌려준다."""
    c = Canvas(d, slide.shapes, ox, oy, f"{TITLES[effect]} (도형)")
    DRAW[(style[1], effect)](c, theme(style, effect))
    return c.g

