"""DRAM 기술 경쟁력 제고를 위한 차별화 준비: 현실 파악 1장, 시각화 시안 3종 (v1, 2026-10-14).

왼쪽 절반 = 위기(Tech node shrink 장벽과 후발 주자의 추격), 오른쪽 절반 = 기회(메모리를 읽는 기계: AI 추론의 병목은 메모리).
같은 내용을 세 가지 시각 언어로 그린다. 사용자가 하나를 고른다.
  시안 A 흐름형     왼쪽 로그 축 노드 추이 / 오른쪽 선순환 → 드라이버 3종 막대 → 총수요 산술식 → 알고리즘 계단 · 공급 시계
  시안 B 선순환형   왼쪽 선형 축 + 2013 이후 확대 / 오른쪽 선순환 고리(가속기 이미지) + 총수요 빅넘버 + 늘리는 힘 vs 줄이는 힘
  시안 C 실물형     왼쪽 로그 축 + 노드 전환 점 줄 / 오른쪽 요청 1건 KV vs HBM 288GB(넘침) · 토큰 추이 · 알고리즘 계단 · 메모리 계층 이미지

규율: samsung-memory-ppt-design-skill v2.2(11절 시각 우선, 11.J 근거 사슬, 20 × 11.25in, 본문 18pt 이상 · 출처 15pt, em-dash 금지).
색: 사용자 지시(2026-10-14 "블루톤만 고집하지 말고 다양하게")로 의미별 색을 쓴다. 위기 = 주황, 기회 = 청록, 알고리즘 = 보라, 넘침 · 격차 = 빨강.
기업은 로고(assets/logos, CXMT는 cxmt.com 공식 SVG 렌더), 부품은 3D 렌더(assets/survival/render_*.png).
근거: sources/articles/dram-shrink-wall-and-ai-memory-bottleneck-2026-10.md, 차트: generate_dram_reality_charts.py
원고: outputs/presentation/dram-reality-outline.md
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(ROOT, ".claude", "skills", "samsung-memory-ppt-design-skill", "scripts"))

from deck_helpers import BLUE, BLUE_T2, CW, GRAY, GRAY_2, INK, LINE, MX, PALE, RIGHT, WHITE, Deck  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402

ASSETS = os.path.join(HERE, "..", "assets")
FIG = os.path.join(ASSETS, "dram-reality")
OUT = os.environ.get("OUT_PATH") or os.path.join(HERE, "..", "dram-reality-options.pptx")

d = Deck("DRAM 기술 경쟁력 제고 · 차별화 준비", total=3,
         logos_dir=os.path.join(ASSETS, "logos"), photos_dir=os.path.join(ASSETS, "photos"),
         renders_dir=os.path.join(ASSETS, "survival"))
tb, rect, label_box = d.tb, d.rect, d.label_box
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
MID = MSO_ANCHOR.MIDDLE
RR = MSO_SHAPE.ROUNDED_RECTANGLE

ORANGE = RGBColor(0xE8, 0x59, 0x0C)
ORANGE_BG = RGBColor(0xFD, 0xEF, 0xE6)
TEAL = RGBColor(0x0E, 0x93, 0x84)
TEAL_BG = RGBColor(0xE6, 0xF3, 0xF1)
PURPLE = RGBColor(0x69, 0x41, 0xC6)
PURPLE_BG = RGBColor(0xF0, 0xEB, 0xFA)
RED = RGBColor(0xD9, 0x30, 0x25)
NAVY = RGBColor(0x1F, 0x2A, 0x5A)

# 좌우 절반
LX, LW = MX, 8.76            # 0.79 ~ 9.55
RX, RW = 10.25, RIGHT - 10.25  # 10.25 ~ 19.21
RE = RIGHT

TITLE = ("DRAM은 Shrink 장벽에 막혀 후발 주자의 추격을 받고 있지만,\n"
         "AI 추론의 병목이 연산에서 메모리로 옮겨 오며 차별화의 기회가 열리고 있다")
BAND = "Shrink만으로는 격차를 지킬 수 없다. 차별화는 AI 추론의 메모리 병목에서 찾는다"

NOTES_COMMON = (
    "[현실 파악] 왼쪽은 위기, 오른쪽은 기회입니다.\n"
    "위기: 업계 선두의 주력 노드는 2000~2011년 연 15%씩 줄었지만 2013~2025년은 연 6%로 느려졌습니다. "
    "새 노드가 1년마다 나오던 것이 1.6년마다로 늘었고, 6F² 셀은 1d까지가 한계로 보고됩니다(TechInsights). "
    "그 사이 CXMT는 EUV 없이 DUV 다중 패터닝으로 16nm(1z급) DDR5를 양산하고, 2026년 말 캐파가 Micron의 93%(35만 vs 37.5만 wpm)에 이릅니다. "
    "노드 격차는 약 3세대(1a · 1b · 1c)지만 세대 간 간격이 길어져 시간으로 버는 격차가 줄어듭니다.\n"
    "기회: 추론은 토큰 하나를 만들 때마다 가중치와 KV 캐시를 읽는 '메모리를 읽는 기계'입니다. "
    "알고리즘(모델 · 문맥 · 추론 SW)은 더 큰 모델, 더 긴 문맥, 더 많은 토큰으로 메모리에 고용량 · 고대역폭을 요구합니다. "
    "산술: 토큰당 KV 320 KiB(Llama 3.1 70B, BF16) × Google 월 3,200조 토큰 ≈ 월 1 ZB의 KV 생성. "
    "토큰이 HBM에 60초~5분 머문다고 가정하면 상주 KV는 24~121 PB, HBM 288GB 가속기 8만~42만 개 분량입니다(가정 의존, 규모 감각용). "
    "알고리즘은 GQA(÷8), MLA, FP8/FP4 KV, PagedAttention(낭비 60~80% → 4% 미만), prefix caching(재사용 86%)으로 덜 쓰는 방법을 만들지만, "
    "같은 기간 토큰은 2년에 330배 늘었고, 메모리 공급은 신규 팹의 실질 기여가 2028년이라 한동안 병목이 이어집니다(TrendForce 2026-07: 2026 충족률 -1~-2%)."
)


# =============================================================== 공통 도구
def head(s, x, w, no, text, color, y=2.34):
    dd = 0.48
    rect(s, x, y, dd, dd, fill=color, shape=MSO_SHAPE.OVAL)
    tb(s, x, y, dd, dd, [(str(no), 20, True, WHITE)], align=C, anchor=MID)
    tb(s, x + dd + 0.16, y - 0.04, w - dd - 0.16, dd + 0.08, [(text, 24, True, INK)], anchor=MID)


def sub(s, x, y, w, text, color, size=18):
    tb(s, x, y, w, 0.34, [(text, size, True, color)], anchor=MID)


def pill(s, x, y, w, h, text, color, fill=WHITE, size=16, bold=True):
    rect(s, x, y, w, h, fill=fill, line=color, lw=1.25, shape=RR)
    tb(s, x, y, w, h, [(text, size, bold, color)], align=C, anchor=MID)


def img(s, name, x, y, bw, bh):
    return d.fit(s, os.path.join(FIG, name), x, y, bw, bh)


def part(s, name, x, y, bw, bh):
    return d.fit(s, d.part(name), x, y, bw, bh)


def logo(s, name, x, y, h):
    w, _ = d.img(s, d.logo(name), x, y, h=h)
    return w


def frame(s, variant):
    """세로 구분선 + 밴드 + 머리글 2개."""
    rect(s, 9.86, 2.40, 0.02, 7.02, fill=LINE)
    head(s, LX, LW, 1, "위기: Shrink 장벽과 후발 주자의 추격", ORANGE)
    head(s, RX, RW, 2, "기회: AI 추론의 병목은 연산이 아닌 메모리", TEAL)
    d.band(s, 9.60, 0.80, "현실", BAND, size=24)


def footer(s, text, tag):
    tb(s, MX, 10.50, 16.9, 0.62, [(text, 15, False, GRAY)], spacing=1.05)
    tb(s, RIGHT - 1.6, 10.50, 1.6, 0.34, [(f"시안 {tag}", 16, True, GRAY)], align=R)


def kpis(s, y=8.44, h=0.98):
    """왼쪽 아래 근거 빅넘버 3칸."""
    gap = 0.18
    ws = [2.60, 2.60, LW - 5.20 - 2 * gap]
    items = [
        ("-15% → -6%", "연간 노드 축소 속도\n2000~11 → 2013~25", ORANGE, None),
        ("1년 → 1.6년", "새 노드가 나오는 간격\n2000~11 → 2013~25", ORANGE, None),
        ("93%", "2026 말 캐파, Micron 대비\n35만 vs 37.5만 wpm", INK, "cxmt"),
    ]
    x = LX
    for i, (big, cap, col, lg) in enumerate(items):
        w = ws[i]
        if i:
            x += ws[i - 1] + gap
        rect(s, x, y, w, h, fill=WHITE, line=LINE, lw=1.0, shape=RR)
        bx = x + 0.16
        if lg:
            lw_ = logo(s, lg, bx, y + 0.13, 0.27)
            bx += lw_ + 0.14
        tb(s, bx, y + 0.04, x + w - bx - 0.1, 0.46, [(big, 26, True, col)], anchor=MID)
        tb(s, x + 0.16, y + 0.50, w - 0.24, 0.46, [(ln, 15, False, GRAY) for ln in cap.split("\n")], spacing=1.0)


def supply_timeline(s, x, y, w):
    """공급 시계: 2026 충족률 → 2H27 램프 → 2028 실질 기여."""
    pts = [("2026", "부족 1~2%"), ("2H27", "신규 팹 램프"), ("2028", "실질 출하")]
    step = w / 3
    rect(s, x + step / 2, y + 0.20, step * 2, 0.04, fill=GRAY_2)
    for i, (yr, txt) in enumerate(pts):
        cx = x + step * i + step / 2
        col = RED if i == 0 else (GRAY if i == 1 else TEAL)
        rect(s, cx - 0.13, y + 0.09, 0.26, 0.26, fill=col, shape=MSO_SHAPE.OVAL)
        tb(s, cx - step / 2, y + 0.38, step, 0.30, [(yr, 18, True, col)], align=C)
        tb(s, cx - step / 2 - 0.05, y + 0.68, step + 0.1, 0.30, [(txt, 15, False, INK)], align=C)


def loop_strip(s, x, y, w, h=0.80):
    """알고리즘 ⇄ 메모리 선순환 띠."""
    bw = 3.10
    label_box(s, x, y, bw, h, [("알고리즘", 20, True, PURPLE), ("모델 · 문맥 · 추론 SW", 17, False, INK)],
              fill=PURPLE_BG, line=PURPLE, lw=1.25)
    label_box(s, x + w - bw, y, bw, h, [("메모리", 20, True, TEAL), ("고용량 · 고대역폭", 17, False, INK)],
              fill=TEAL_BG, line=TEAL, lw=1.25)
    mx, mw = x + bw + 0.12, w - 2 * bw - 0.24
    rect(s, mx, y + 0.10, mw, 0.20, fill=PURPLE, shape=MSO_SHAPE.RIGHT_ARROW)
    rect(s, mx, y + h - 0.30, mw, 0.20, fill=TEAL, shape=MSO_SHAPE.LEFT_ARROW)
    tb(s, mx, y + 0.27, mw, 0.27, [("더 많이 요구 · 더 크게 허용", 15, False, GRAY)], align=C, anchor=MID)


SRC_LEFT = ("노드: ITRS · TechInsights · 각 사 발표(선두 주력 양산 근사) · 2026 위치: Fab Tracker Q2'26 · Digitimes · "
            "SemiAnalysis")
SRC_RIGHT = ("KV: Llama 3.1 · DeepSeek-V3 설정 산식 · 토큰: Google I/O · Paged: SOSP'23 · 공급: TrendForce 2026-07 · "
             "산술은 가정 의존 · 부품은 3D 렌더")


# =============================================================== 시안 A: 흐름형
def slide_a():
    s = d.slide(1, "현실 파악 · 시안 A 흐름형", TITLE)
    frame(s, "A")
    # 왼쪽
    img(s, "node_log.png", LX, 2.98, LW, 5.30)
    kpis(s)
    # 오른쪽 1: 선순환 띠
    loop_strip(s, RX, 2.98, RW)
    # 오른쪽 2: 드라이버 3종
    img(s, "drivers.png", RX, 3.92, RW, 2.00)
    # 오른쪽 3: 총수요 산술식
    tb(s, RX, 6.02, RW, 0.30, [[("메모리 총수요 산술", 18, True, INK),
                                ("   KV/토큰 × 월간 토큰 → 상주량 (가정: HBM 체류 60초~5분)", 15, False, GRAY)]], anchor=MID)
    y, h = 6.38, 0.86
    boxes = [
        (1.80, "320 KiB", "KV / 토큰 (70B)", PURPLE, PURPLE_BG),
        (0.34, "×", None, GRAY, None),
        (1.80, "3,200조", "Google 월간 토큰", TEAL, TEAL_BG),
        (0.34, "=", None, GRAY, None),
        (1.62, "≈1 ZB", "월간 KV 생성", ORANGE, ORANGE_BG),
        (0.34, "→", None, GRAY, None),
        (2.40, "8만~42만 개", "HBM 288GB 가속기분", RED, WHITE),
    ]
    x = RX
    for bw, big, cap, col, bg in boxes:
        if cap is None:
            tb(s, x, y, bw, h, [(big, 28, True, col)], align=C, anchor=MID)
        else:
            rect(s, x, y, bw, h, fill=bg, line=col, lw=1.5 if col == RED else 1.0, shape=RR)
            tb(s, x, y + 0.06, bw, 0.44, [(big, 24, True, col)], align=C, anchor=MID)
            tb(s, x, y + 0.48, bw, 0.32, [(cap, 15, False, INK)], align=C, anchor=MID)
        x += bw
    # 오른쪽 4: 알고리즘 계단 + 공급 시계
    hw = 4.30
    sub(s, RX, 7.40, hw, "알고리즘이 줄인다 (KiB / 토큰)", PURPLE)
    img(s, "kv_stair4.png", RX, 7.78, hw, 1.30)
    tb(s, RX, 9.12, hw + 0.3, 0.30, [("Paged 낭비 70%→4% · Prefix 재사용 86%", 15, False, PURPLE)], anchor=MID)
    sx = RX + hw + 0.36
    sw = RE - sx
    sub(s, sx, 7.40, sw, "메모리는 바로 늘릴 수 없다", RED)
    part(s, "dram", sx, 7.80, 1.55, 0.62)
    tb(s, sx + 1.65, 7.80, sw - 1.65, 0.62, [("팹 건설 · 장비 · HBM 웨이퍼 소모", 15, False, GRAY)], anchor=MID)
    supply_timeline(s, sx, 8.45, sw)
    footer(s, "출처: " + SRC_LEFT + " / " + SRC_RIGHT, "A")
    d.notes(s, NOTES_COMMON + "\n[시안 A] 왼쪽 로그 축 + 두 시기 회귀 띠, 오른쪽은 선순환 → 수요 드라이버 → 총수요 산술 → 줄이는 힘과 공급 한계 순서로 읽습니다.")


# =============================================================== 시안 B: 선순환형
def slide_b():
    s = d.slide(2, "현실 파악 · 시안 B 선순환형", TITLE)
    frame(s, "B")
    img(s, "node_linear.png", LX, 2.98, LW, 5.30)
    kpis(s)
    # 오른쪽 위: 선순환 고리
    ring = 2.28
    rx, ry = RX + 0.41, 3.36
    tb(s, RX, 2.98, 3.10, 0.36, [("알고리즘", 20, True, PURPLE)], align=C, anchor=MID)
    img(s, "flywheel.png", rx, ry, ring, ring)
    cx, cy = rx + ring / 2, ry + ring / 2
    part(s, "hbm", cx - 0.62, cy - 0.48, 1.24, 0.96)
    tb(s, RX, ry + ring + 0.02, 3.10, 0.36, [("메모리", 20, True, TEAL)], align=C, anchor=MID)
    # 고리 오른쪽: 요구 3종 + 총수요
    qx = RX + 3.45
    qw = RE - qx
    tb(s, qx, 3.00, qw, 0.34, [("알고리즘이 메모리에 요구하는 것", 18, True, PURPLE)], anchor=MID)
    for i, (k, v) in enumerate([("더 큰 모델", "KV/토큰 128 → 504 KiB"),
                                ("더 긴 문맥", "요청 1건 2.7 → 328 GB"),
                                ("더 많은 토큰", "Google 월 9.7조 → 3,200조")]):
        yy = 3.44 + i * 0.42
        rect(s, qx, yy + 0.08, 0.08, 0.24, fill=ORANGE)
        tb(s, qx + 0.20, yy, 1.70, 0.40, [(k, 18, True, INK)], anchor=MID)
        tb(s, qx + 1.90, yy, qw - 1.90, 0.40, [(v, 18, False, INK)], anchor=MID)
    # 총수요 빅넘버
    by = 4.86
    rect(s, qx, by, qw, 1.22, fill=ORANGE_BG, line=ORANGE, lw=1.25, shape=RR)
    tb(s, qx + 0.20, by + 0.06, qw - 0.4, 0.34, [("320 KiB × 월 3,200조 토큰 ≈ 월 1 ZB KV", 16, True, INK)], anchor=MID)
    tb(s, qx + 0.20, by + 0.42, qw - 0.4, 0.50, [[("HBM 가속기 ", 20, True, INK), ("8만~42만 개", 30, True, RED), (" 분량", 20, True, INK)]], anchor=MID)
    tb(s, qx + 0.20, by + 0.90, qw - 0.4, 0.28, [("산술 · 가정: 288GB 가속기, HBM 체류 60초~5분", 15, False, GRAY)], anchor=MID)
    # 오른쪽 아래: 늘리는 힘 vs 줄이는 힘
    sub(s, RX, 6.30, RW, "줄이는 힘보다 늘리는 힘이 크다 (2024~2026, 요인별 배수)", INK)
    img(s, "tug.png", RX, 6.66, RW, 2.35)
    rect(s, RX, 9.06, RW, 0.36, fill=PALE, shape=RR)
    tb(s, RX + 0.16, 9.06, RW - 0.3, 0.36, [[("공급 한계  ", 16, True, RED),
                                              ("2026 충족률 -1~-2%, 신규 팹 실질 기여는 2028년", 16, False, INK)]], anchor=MID)
    footer(s, "출처: " + SRC_LEFT + " / " + SRC_RIGHT + " · 배수는 요인별(곱하지 않음)", "B")
    d.notes(s, NOTES_COMMON + "\n[시안 B] 왼쪽 선형 축으로 2013년 이후의 평탄화를 그대로 보여 주고, 확대 상자에서 각 사 위치를 봅니다. "
            "오른쪽은 알고리즘과 메모리의 선순환 고리, 총수요 빅넘버, 그리고 줄이는 힘(보라)과 늘리는 힘(주황)의 배수 비교입니다. "
            "배수는 요인별 정의가 달라 서로 곱하지 않습니다(GQA = 헤드 수, MLA = 405B GQA 대비, FP4 = BF16 대비, Paged = 유효 메모리 30% → 96%, "
            "모델 = 8B → 405B KV/토큰, 문맥 = 8K → 1M, 토큰 = Google 2024-05 → 2026-05).")


# =============================================================== 시안 C: 실물형
def slide_c():
    s = d.slide(3, "현실 파악 · 시안 C 실물형", TITLE)
    frame(s, "C")
    img(s, "node_cadence.png", LX, 2.98, LW, 5.30)
    kpis(s)
    # 오른쪽 위 왼쪽: 요청 1건 vs HBM
    tw = 4.55
    img(s, "tank.png", RX, 3.00, tw, 3.20)
    part(s, "hbm", RX + 0.35, 4.14, 1.45, 0.95)
    # 오른쪽 위 오른쪽: 토큰 추이 + 산술
    kx = RX + tw + 0.20
    kw = RE - kx
    img(s, "tokens.png", kx, 3.00, kw, 2.50)
    rect(s, kx, 5.56, kw, 0.64, fill=ORANGE_BG, line=ORANGE, lw=1.0, shape=RR)
    tb(s, kx, 5.56, kw, 0.64, [[("× 320 KiB/토큰 = ", 17, True, INK), ("월 ≈1 ZB KV", 22, True, ORANGE)]], align=C, anchor=MID)
    # 오른쪽 아래 왼쪽: 알고리즘 계단
    hw = 4.35
    sub(s, RX, 6.36, hw, "알고리즘이 줄인다 (KiB / 토큰)", PURPLE)
    img(s, "kv_stair.png", RX, 6.74, hw, 2.30)
    tb(s, RX, 9.10, hw + 0.25, 0.32, [("Paged 낭비 70%→4% · Prefix 재사용 86%", 15, False, PURPLE)], anchor=MID)
    # 오른쪽 아래 오른쪽: 메모리 계층 + 공급 시계
    sx = RX + hw + 0.30
    sw = RE - sx
    sub(s, sx, 6.36, sw, "넘치면 아래 계층으로, 공급은 2028년", RED)
    names = [("hbm", "HBM"), ("dram", "DRAM"), ("server", "SSD · 서버")]
    cw, gap = 1.12, (sw - 3 * 1.12) / 2
    for i, (nm, lab) in enumerate(names):
        x = sx + i * (cw + gap)
        part(s, nm, x, 6.80, cw, 0.80)
        tb(s, x - 0.2, 7.62, cw + 0.4, 0.30, [(lab, 16, True, INK)], align=C, anchor=MID)
        if i < 2:
            d.arrow_r(s, x + cw + 0.06, 7.07, gap - 0.12, 0.24, fill=RED if i == 0 else GRAY_2)
    supply_timeline(s, sx, 8.08, sw)
    footer(s, "출처: " + SRC_LEFT + " / " + SRC_RIGHT + " · KV 계층화: Google Next '26", "C")
    d.notes(s, NOTES_COMMON + "\n[시안 C] 왼쪽 아래 점 줄이 새 노드의 간격(촘촘 → 성김)을 보여 줍니다. "
            "오른쪽은 실물 감각: 요청 1건의 KV가 1M 문맥이면 가속기 1개의 HBM(288GB)을 넘고, 넘친 KV는 DRAM · SSD 계층으로 내려갑니다"
            "(Google Next '26 'KV Cache storage tiering across RAM, Local SSD'). 토큰은 2년에 330배, 알고리즘은 토큰당 KV를 MHA 대비 최대 1/37로 줄였습니다.")


slide_a()
slide_b()
slide_c()
print("saved", d.save(OUT))
