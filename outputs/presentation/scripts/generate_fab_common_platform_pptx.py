"""DRAM/NAND Fab 통합 운영: Bonding 시대에 맞는 Common Platform 구성 (1장, 시안 6종) v2.1 (2026-10-13).

v2.1: 같은 3종을 다색 버전으로 한 번 더(시안 4~6). 색 의미: DRAM = Blue, NAND = Teal, 효과 ① Blue ② Teal ③ Purple,
      개선 · 운영 = Amber, 지붕 · 기초 = Navy. 그림은 c{1,2,3}_*.png. (단일 액센트 규칙은 사용자 요청으로 시안 4~6에서만 해제)

v2.0: v1.0 시안 C(건축형)의 형식을 세 장 모두에 쓰고, 기둥 머리 그림을 크게(5.30 x 1.80in) 키워 그림 스타일만 달리했다.
  지붕(결론) · 기둥 3개(머리 = 그림 + 효과 이름, 몸통 = 항목 벽돌 3장) · Bonding 기초(Common Platform) · 맨 아래 운영 제언(두 조직의 Bonding)
  시안 1 아이소메트릭  3D 느낌의 Fab 바닥 · 설비 블록
  시안 2 메커니즘      와플 격자 · 수요 곡선 · 막대 (방향만 보이는 개념도)
  시안 3 Fab 도면      위에서 본 Bay · OHT 레일 · 웨이퍼 Lot 흐름 · 설비 모델 매트릭스
그림은 generate_fab_common_platform_figs.py가 만든다(assets/fab-common-platform/v{1,2,3}_*.png).

규율: samsung-memory-ppt-design-skill v2.2 (20 x 11.25in, 본문 18pt 이상 · 그림 라벨 16pt · 출처 15pt, em-dash 금지,
액센트는 Samsung Blue 하나, 수치를 지어내지 않는다).
출력: outputs/presentation/fab-common-platform-options.pptx
재생성: .venv/bin/python outputs/presentation/scripts/generate_fab_common_platform_figs.py
        .venv/bin/python outputs/presentation/scripts/generate_fab_common_platform_pptx.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(ROOT, ".claude", "skills", "samsung-memory-ppt-design-skill", "scripts"))

from deck_helpers import (BLUE, BLUE_T1, BLUE_T2, CW, GRAY, GRAY_2, INK, MX, PALE,  # noqa: E402
                          RIGHT, TINT, WHITE, Deck)
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

ASSETS = os.path.join(ROOT, "outputs", "presentation", "assets", "fab-common-platform")
OUT = os.environ.get("OUT_PATH", os.path.join(ROOT, "outputs", "presentation", "fab-common-platform-options.pptx"))

TITLE = "DRAM/NAND Fab 통합 운영:\nBonding 시대에 맞는 Common Platform 구성"
FOOT = "주: 효과 그림은 방향을 보여 주는 개념도이며 수치는 표기하지 않음. 정량 효과는 Fab별 검증 후 확정 [사내 확인]"

# (번호, 국문, 영문, [(굵게, 나머지), ...])
EFFECTS = [
    (1, "공간 효율성", "Space Efficiency", [
        ("설비 밀집도 향상", ", Dead Space 최소화"),
        ("동일 모델 군집화", " 시너지"),
        ("단위 면적당 생산성", " 극대화"),
    ]),
    (2, "시황 변화 대응력", "Market Agility", [
        ("DRAM/NAND 설비 전환", " 속도 증가"),
        ("생산 Bottleneck 해소", "로 생산성 증가"),
        ("유휴 설비 최소화", ", 투자 회수 기간 단축"),
    ]),
    (3, "설비 통일성·선단화", "Equipment Unif.", [
        ("P-Mix 관점 ", "투자·전환 Loss 감소"),
        ("생산성 향상", ", 수율 안정화"),
        ("강력한 구매 협상력", " 확보"),
    ]),
]


# ------------------------------------------------------------------ 공통 도구
def line(s, x1, y1, x2, y2, color=GRAY_2, w=1.5, head=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def circ_no(d, s, x, y, n, dd=0.42, size=18, fill=BLUE):
    d.rect(s, x, y, dd, dd, fill=fill, shape=MSO_SHAPE.OVAL)
    d.tb(s, x, y, dd, dd, [(str(n), size, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def people(d, s, x, y, h, n, color, gap=0.10):
    for i in range(n):
        d.person(s, x + i * (0.62 * h + gap), y, h, color)
    return n * (0.62 * h + gap) - gap


# ================================================================== 색 테마
# MONO: Samsung Blue 단일 액센트(시안 1~3). COLOR: 다색(시안 4~6, 사용자 요청 2026-10-13).
# COLOR의 색 의미는 그림과 같다: DRAM = Blue, NAND = Teal, 효과 ① Blue ② Teal ③ Purple, 개선 · 운영 = Amber, 지붕 · 기초 = Navy.
C_BLUE, C_TEAL, C_PURPLE = RGBColor(0x3C, 0x4F, 0xA5), RGBColor(0x1E, 0x8C, 0x8C), RGBColor(0x6E, 0x56, 0xB0)
C_AMBER, C_AMBER_DK, C_AMBER_LT = RGBColor(0xE3, 0x9B, 0x2D), RGBColor(0xB8, 0x6E, 0x00), RGBColor(0xF6, 0xC4, 0x6B)
C_NAVY = RGBColor(0x26, 0x33, 0x6E)
MONO_T = dict(roof=BLUE, roof_sub=BLUE_T2, roof_hot=WHITE, pillars=[(BLUE, BLUE_T2, TINT)] * 3, pads=[BLUE_T1], found=BLUE,
              orgD=(PALE, GRAY_2, GRAY_2), orgN=(PALE, GRAY_2, GRAY_2), bpads=[BLUE_T1], bond=BLUE, arrow=BLUE, box=BLUE,
              chips=[BLUE] * 3, up=BLUE_T1, up_txt=BLUE, tag=BLUE)
COLOR_T = dict(roof=C_NAVY, roof_sub=RGBColor(0xAE, 0xB8, 0xE0), roof_hot=C_AMBER_LT,
               pillars=[(C_BLUE, RGBColor(0xC0, 0xCA, 0xEC), RGBColor(0xF0, 0xF2, 0xF8)),
                        (C_TEAL, RGBColor(0xA9, 0xDC, 0xD7), RGBColor(0xE8, 0xF5, 0xF4)),
                        (C_PURPLE, RGBColor(0xD3, 0xCA, 0xEE), RGBColor(0xF1, 0xEE, 0xF9))],
               pads=[C_BLUE, C_TEAL], found=C_NAVY,
               orgD=(RGBColor(0xE9, 0xEC, 0xF8), C_BLUE, C_BLUE), orgN=(RGBColor(0xE3, 0xF3, 0xF1), C_TEAL, C_TEAL),
               bpads=[C_BLUE, C_TEAL], bond=C_NAVY, arrow=C_AMBER, box=C_NAVY, chips=[C_BLUE, C_TEAL, C_PURPLE],
               up=C_AMBER, up_txt=C_AMBER_DK, tag=C_NAVY)


# ================================================================== 공통 골격 (v1.0 시안 C 형식)
STYLES = [
    ("시안 1 · 아이소메트릭", "v1", "기둥 그림은 3D 느낌의 Fab 바닥입니다. 1 공간: 분리 Fab은 DRAM과 NAND 사이에 벽이 있고 모델이 제각각이며 빈 칸이 생기지만, "
     "통합 Fab은 같은 모델을 줄지어 놓아 한 줄이 통째로 남습니다. 2 대응력: 공용 설비가 DRAM과 NAND 사이를 오가며, 수요 곡선이 엇갈리는 대로 설비 배분이 따라갑니다. "
     "3 통일성: 높이와 크기가 제각각인 모델을 같은 모델로 모으면 모델당 발주가 한 단에서 여러 단으로 쌓입니다."),
    ("시안 2 · 메커니즘", "v2", "기둥 그림은 원리를 보이는 개념 차트입니다. 1 공간: 같은 Fab 면적을 칸으로 나누면 분리 Fab의 Dead Space 칸이 통합 후 추가 설비 칸이 됩니다. "
     "2 대응력: DRAM과 NAND의 수요 피크가 엇갈려 합산 피크가 각 피크의 합보다 낮고, 그 차이가 유휴 설비 감소입니다. 설비 전환 시간과 Bottleneck 대기도 줄어듭니다. "
     "3 통일성: 설비 모델 수는 줄고 모델당 발주 물량은 늘어 구매 협상력이 생깁니다. 수치는 Fab별 검증 후 채웁니다."),
    ("시안 3 · Fab 도면", "v3", "기둥 그림은 엔지니어링 도면 언어입니다. 1 공간: 위에서 본 Fab에서 분리 시 OHT 루프 2개와 제각각인 설비 Footprint, 빈 Bay가 생기지만, "
     "통합 시 OHT 루프 하나에 같은 모델 줄 배치로 Bay 하나가 남습니다. 2 대응력: DRAM과 NAND 웨이퍼 Lot이 공용 BEOL 설비 Pool의 어느 설비로든 들어가고, "
     "결과물은 Bonding으로 이어집니다. 3 통일성: 공정 단계마다 DRAM 모델 둘, NAND 모델 둘이던 것을 공용 모델 하나로 모아 레시피 · 부품 · 물량을 공유합니다. 모델 코드는 예시입니다."),
]
FIGS = ["space", "agility", "unif"]


def roof_pillars_base(d, s, style, T):
    roof = d.rect(s, MX, 2.28, CW, 0.70, fill=T["roof"], shape=MSO_SHAPE.TRAPEZOID)
    roof.adjustments[0] = 0.55
    d.tb(s, MX + 1.6, 2.28, CW - 3.2, 0.70,
         [[("종합 반도체 경쟁력  ", 18, False, T["roof_sub"]), ("기술·생산 경쟁력 + 효율·유연성 + ", 24, True, WHITE),
           ("Buying Power", 24, True, T["roof_hot"])]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    px0, pgap = MX + 0.30, 0.42
    pw = (CW - 0.60 - 2 * pgap) / 3
    for k, (n, ko, en, rows) in enumerate(EFFECTS):
        x = px0 + k * (pw + pgap)
        hue, hue_lt, bg = T["pillars"][k]
        d.rect(s, x, 3.08, pw, 2.44, fill=bg)
        d.rect(s, x, 3.08, pw, 0.06, fill=hue)
        circ_no(d, s, x + 0.18, 3.22, n, dd=0.40, size=18, fill=hue)
        d.tb(s, x + 0.70, 3.18, pw - 0.86, 0.48, [[(ko, 24, True, INK), ("  " + en, 16, False, hue)]], anchor=MSO_ANCHOR.MIDDLE)
        d.img(s, os.path.join(ASSETS, f"{style}_{FIGS[k]}.png"), x + (pw - 5.30) / 2, 3.68, w=5.30)
        for i, (b, rest) in enumerate(rows):
            yy = 5.60 + i * 0.56
            d.rect(s, x, yy, pw, 0.50, fill=WHITE, line=hue_lt, lw=1.25)
            d.tb(s, x + 0.22, yy, pw - 0.36, 0.50, [[(b, 18, True, hue), (rest, 18, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)

    fx, fw = MX + 0.30, CW - 0.60
    n_pad = 44
    for i in range(n_pad):
        d.rect(s, fx + 0.10 + i * (fw - 0.32) / (n_pad - 1), 7.34, 0.12, 0.08, fill=T["pads"][i % len(T["pads"])])
    d.rect(s, fx, 7.44, fw, 0.56, fill=T["found"])
    d.tb(s, fx + 0.30, 7.44, fw - 0.60, 0.56,
         [[("Common Platform   ", 22, True, WHITE), ("DRAM·NAND 공용 BEOL · Bonding 공정 · 설비 Pool", 20, False, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def org_row(d, s, T):
    """운영 제언: DRAM · NAND BEOL 조직을 웨이퍼처럼 Bonding해 통합조직으로, 통합조직이 기초를 받쳐 운영."""
    d.rect(s, MX, 8.38, 1.70, 0.46, fill=WHITE, line=T["tag"], lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    d.tb(s, MX, 8.38, 1.70, 0.46, [("운영 제언", 20, True, T["tag"])], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    d.tb(s, MX, 8.98, 3.70, 1.00, [("메모리 BEOL", 26, True, INK), ("통합조직 구성", 26, True, INK)], spacing=1.05)

    wx, ww = 4.90, 4.40
    for lab, yy, (fc, lc, pc) in (("DRAM BEOL 조직", 8.36, T["orgD"]), ("NAND BEOL 조직", 9.54, T["orgN"])):
        d.rect(s, wx, yy, ww, 0.64, fill=fc, line=lc, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        people(d, s, wx + 0.22, yy + 0.12, 0.42, 3, pc, gap=0.08)
        d.tb(s, wx + 1.40, yy, ww - 1.50, 0.64, [(lab, 18, True, INK)], anchor=MSO_ANCHOR.MIDDLE)
    for i in range(12):
        d.rect(s, wx + 0.30 + i * (ww - 0.72) / 11, 9.10, 0.12, 0.34, fill=T["bpads"][i % len(T["bpads"])])
    d.tb(s, wx + ww + 0.12, 9.08, 1.20, 0.38, [("Bonding", 18, True, T["bond"])], anchor=MSO_ANCHOR.MIDDLE)
    d.arrow_r(s, wx + ww + 1.36, 9.09, 0.62, 0.36, fill=T["arrow"])

    bx, bw, by, bh = 11.40, 5.10, 8.36, 1.82
    d.rect(s, bx, by, bw, bh, fill=T["box"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    d.tb(s, bx, by + 0.08, bw, 0.46, [("메모리 BEOL 통합조직", 24, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ppw = 6 * (0.62 * 0.44 + 0.12) - 0.12
    people(d, s, bx + (bw - ppw) / 2, by + 0.60, 0.44, 6, WHITE, gap=0.12)
    cw_, cg = 1.50, 0.12
    cx0 = bx + (bw - 3 * cw_ - 2 * cg) / 2
    for i, lab in enumerate(("표준 Layout", "설비 Pool", "통합 구매")):
        cx = cx0 + i * (cw_ + cg)
        d.rect(s, cx, by + 1.18, cw_, 0.48, fill=T["chips"][i], line=WHITE, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        d.tb(s, cx, by + 1.18, cw_, 0.48, [(lab, 18, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for ax_ in (bx + 1.20, bx + bw - 1.20):
        d.rect(s, ax_ - 0.17, 8.06, 0.34, 0.26, fill=T["up"], shape=MSO_SHAPE.UP_ARROW)
    d.tb(s, bx + 1.50, 8.05, bw - 3.00, 0.30, [("운영", 16, True, T["up_txt"])], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    tx = bx + bw + 0.30
    line(s, bx + bw, by + bh / 2, tx, by + bh / 2, GRAY_2, 1.5)
    d.rect(s, tx, by + 0.34, RIGHT - tx, 1.14, fill=WHITE, line=GRAY_2, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    d.tb(s, tx, by + 0.34, RIGHT - tx, 1.14, [("산하", 16, False, GRAY), ("반도체연구소", 18, True, INK), ("/ MTC", 18, True, INK)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)


def make_slide(d, no, label, style, note, T):
    s = d.slide(no, label, TITLE)
    roof_pillars_base(d, s, style, T)
    org_row(d, s, T)
    d.footer(s, FOOT + (". 설비 모델 코드는 예시" if style[1] == "3" else ""))
    d.notes(s, note + " 지붕은 결론, 기초는 Common Platform, 맨 아래는 DRAM과 NAND의 BEOL 조직을 Bonding해 만든 통합조직이 "
               "반도체연구소 / MTC 산하에서 기초를 운영하는 그림입니다. 통합조직 안의 세 기능(표준 Layout · 설비 Pool · 통합 구매)은 효과 1 · 2 · 3과 짝입니다.")


def main():
    d = Deck("Fab 통합 운영", total=2 * len(STYLES))
    for no, (label, style, note) in enumerate(STYLES, 1):
        make_slide(d, no, label, style, note, MONO_T)
    for no, (label, style, note) in enumerate(STYLES, len(STYLES) + 1):
        label = f"시안 {no} · {label.split(' · ')[1]} 컬러"
        note = ("컬러 버전입니다. 색에 뜻이 있습니다: DRAM = Blue, NAND = Teal, 통합 후 같은 색 = 같은 설비 모델, 공용 모델 = Purple, "
                "개선(확보 · 추가 · 협상력) = Amber, 효과 ① Blue ② Teal ③ Purple. " + note)
        make_slide(d, no, label, "c" + style[1], note, COLOR_T)
    # 테마 스타일(p:style)의 effectRef가 그림자를 그리므로 지운다. 채움 · 선은 모두 명시값이다.
    for sl in d.prs.slides:
        for sh in sl.shapes:
            st = sh._element.find(qn("p:style"))
            if st is not None:
                sh._element.remove(st)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    d.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
