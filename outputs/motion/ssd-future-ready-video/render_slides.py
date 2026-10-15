"""덱 1~4장을 영상용 4K 이미지로 렌더 — 클릭 전(b) · 클릭 후(a) 두 상태

  python3 outputs/motion/ssd-future-ready-video/render_slides.py
    → slides/s1a.jpg · s1b.jpg … s4a.jpg (3840 × 2160, 미커밋)

각 장의 클릭 애니메이션(나타나기 그룹 1개)을 지운 사본이 클릭 전 상태다(4장은 애니메이션 없음 → a만).
글꼴은 웹 보기 이미지와 같이 NanumGothic으로 바꿔 soffice → PDF → pdftoppm.
"""
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

DIR = Path(__file__).resolve().parent
ROOT = DIR.parents[2]
DECK = ROOT / "outputs/presentation/ssd-future-ready-strategy.pptx"
OUT = DIR / "slides"
FONT = b"NanumGothic"
N = 4


def retarget_fonts(src: Path, dst: Path):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml") and (item.filename.startswith("ppt/slides/") or "theme" in item.filename
                                                   or "slideMaster" in item.filename or "slideLayout" in item.filename):
                data = re.sub(rb'(<a:(?:latin|ea|cs) typeface=")[^"]*(")', rb"\g<1>" + FONT + rb"\g<2>", data)
            zout.writestr(item, data)


def strip_click_groups(src: Path, dst: Path):
    """클릭으로 나타나는 도형을 지운 사본 (클릭 전 화면)."""
    prs = Presentation(src)
    for slide in list(prs.slides)[:N]:
        timing = slide._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
        if timing is None:
            continue
        ids = set(re.findall(r'spTgt spid="(\d+)"', etree.tostring(timing).decode()))
        for shape in list(slide.shapes):
            if str(shape.shape_id) in ids:
                shape._element.getparent().remove(shape._element)
        timing.getparent().remove(timing)
    prs.save(dst)


def render(pptx: Path, tag: str, tmp: Path):
    fixed = tmp / f"{tag}.pptx"
    retarget_fonts(pptx, fixed)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", fixed.name], cwd=tmp,
                   check=True, capture_output=True, timeout=300)
    subprocess.run(["pdftoppm", "-f", "1", "-l", str(N), "-jpeg", "-jpegopt", "quality=92",
                    "-scale-to-x", "3840", "-scale-to-y", "2160", tmp / f"{tag}.pdf", tmp / tag], check=True)
    for i in range(1, N + 1):
        page = next(tmp.glob(f"{tag}-{i}.jpg"), None) or next(tmp.glob(f"{tag}-0{i}.jpg"))
        shutil.move(page, OUT / f"s{i}{tag}.jpg")


def main():
    OUT.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        render(DECK, "a", tmp)
        strip_click_groups(DECK, tmp / "before.pptx")
        render(tmp / "before.pptx", "b", tmp)
    (OUT / "s4b.jpg").unlink(missing_ok=True)
    print(sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
